"""Build the 12-slide Huawei ICT submission deck.

Usage:
    python scripts/build_ppt.py

Requires python-pptx and Pillow.  All screenshots come from evidence/.
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
EVIDENCE = ROOT / "evidence"
OUT = ROOT / "submission" / "知径EduPath-华为ICT2026-作品介绍.pptx"

FONT = "微软雅黑"
INK = RGBColor(0x15, 0x21, 0x2C)
INK_2 = RGBColor(0x45, 0x53, 0x60)
MUTED = RGBColor(0x7C, 0x87, 0x92)
FAINT = RGBColor(0xA8, 0xB0, 0xB8)
RED = RGBColor(0xC8, 0x10, 0x2E)
RED_DARK = RGBColor(0x8F, 0x12, 0x24)
RED_SOFT = RGBColor(0xFB, 0xEC, 0xEE)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
TEAL_SOFT = RGBColor(0xE8, 0xF4, 0xF1)
AMBER = RGBColor(0xA8, 0x64, 0x12)
AMBER_SOFT = RGBColor(0xFB, 0xF1, 0xE2)
BLUE = RGBColor(0x24, 0x5C, 0x8A)
BLUE_SOFT = RGBColor(0xEA, 0xF1, 0xF7)
PAPER = RGBColor(0xF7, 0xF4, 0xEE)
WHITE = RGBColor(0xFF, 0xFD, 0xF9)
LINE = RGBColor(0xE2, 0xDB, 0xCF)
DARK = RGBColor(0x10, 0x1F, 0x2B)
DARK_2 = RGBColor(0x17, 0x2A, 0x38)
GRAY_LIGHT = RGBColor(0xD8, 0xD2, 0xC7)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide(bg: RGBColor = PAPER):
    s = prs.slides.add_slide(BLANK)
    shape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg
    shape.line.fill.background()
    shape.shadow.inherit = False
    return s


def rect(s, x, y, w, h, fill, line=None, radius=False, line_width=1.0):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = s.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_width)
    if radius:
        try:
            shape.adjustments[0] = 0.08
        except Exception:
            pass
    shape.shadow.inherit = False
    return shape


def text(s, value, x, y, w, h, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, line=1.2, font=FONT, italic=False, spacing=0):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = value.split("\n")
    for index, item in enumerate(lines):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = line
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(spacing)
        run = paragraph.add_run()
        run.text = item
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        # Force East Asian font.
        from pptx.oxml.ns import qn
        rpr = run._r.get_or_add_rPr()
        ea = rpr.find(qn("a:ea"))
        if ea is None:
            ea = rpr.makeelement(qn("a:ea"), {})
            rpr.append(ea)
        ea.set("typeface", font)
    return box


def picture(s, path, x, y, w, h, border=True):
    path = str(path)
    iw, ih = Image.open(path).size
    target = w / h
    image_ratio = iw / ih
    pic = s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    if image_ratio > target:
        crop = (1 - target / image_ratio) / 2
        pic.crop_left = crop
        pic.crop_right = crop
    elif image_ratio < target:
        crop = (1 - image_ratio / target) / 2
        pic.crop_top = crop
        pic.crop_bottom = crop
    if border:
        pic.line.color.rgb = LINE
        pic.line.width = Pt(1)
    return pic


def header(s, number, section, title, dark=False):
    color = WHITE if dark else INK
    sub = RGBColor(0xB9, 0xC4, 0xCC) if dark else MUTED
    text(s, f"{number} / {section}", 0.7, 0.42, 5.0, 0.3, 10, True, RED if not dark else RGBColor(0xE8, 0x7C, 0x8C))
    text(s, title, 0.7, 0.72, 12.0, 0.75, 27, True, color, line=1.15)
    rect(s, 0.7, 1.52, 0.72, 0.045, RED)


def footer(s, page, dark=False):
    color = RGBColor(0x8A, 0x99, 0xA4) if dark else FAINT
    text(s, "知径 EduPath · 华为 ICT 2026 · 赛题 2", 0.7, 7.12, 5.0, 0.24, 8, False, color)
    text(s, f"{page:02d}", 12.0, 7.12, 0.63, 0.24, 8, False, color, PP_ALIGN.RIGHT)


def pill(s, label, x, y, w, fill, color, size=9):
    rect(s, x, y, w, 0.3, fill, None, radius=True)
    text(s, label, x, y + 0.03, w, 0.22, size, True, color, PP_ALIGN.CENTER)


# ------------------------------------------------------------------ slide 1
s = slide(DARK)
# Right image panel
picture(s, EVIDENCE / "demo-home.png", 7.15, 0.55, 5.75, 6.35, border=False)
rect(s, 7.15, 0.55, 5.75, 6.35, None, RGBColor(0x3A, 0x4C, 0x5A), False, 1)
rect(s, 0, 0, 0.16, 7.5, RED)
rect(s, 0.85, 0.75, 0.52, 0.52, RED, None, True)
text(s, "径", 0.85, 0.82, 0.52, 0.4, 22, True, WHITE, PP_ALIGN.CENTER)
text(s, "第十一届华为 ICT 大赛 · 中国创新赛 · 赛题 2", 1.55, 0.78, 5.4, 0.3, 10.5, True, RGBColor(0xD9, 0xC9, 0xC2))
text(s, "知径 EduPath", 0.85, 1.55, 6.0, 0.9, 44, True, WHITE)
text(s, "基于华为云 AgentArts 的\n课程学习多智能体", 0.85, 2.55, 5.8, 1.25, 23, False, RGBColor(0xF0, 0xEA, 0xDF), line=1.35)
rect(s, 0.88, 4.05, 0.5, 0.035, RED)
text(s, "让每一个薄弱点，都有证据、有路径、有反馈。", 0.85, 4.3, 5.6, 0.35, 13, True, RGBColor(0x8F, 0xC9, 0xC1))
text(s, "作品介绍 · 初赛版\n2026 年 10 月", 0.85, 5.9, 4.0, 0.7, 10.5, False, RGBColor(0x93, 0xA2, 0xAD), line=1.5)

# ------------------------------------------------------------------ slide 2
s = slide()
header(s, "01", "问题", "不是缺资料，是缺“按我薄弱点组织、且有依据”的资料")
cards = [
    ("学生", "“我知道没学好，但不知道弱在哪、下一步学什么。”", ["通用 AI 给泛化内容", "练完没有反馈", "路径一次推荐就不再变"]),
    ("教师", "“一个班几十人，我只能按平均进度讲。”", ["重复备课", "难以按薄弱点分层", "缺少可解释的学情依据"]),
    ("通用 AI", "“内容很多，但我不敢信。”", ["脱离教材", "无法核验来源", "可能直接给答案"]),
]
for index, (title_, quote, bullets) in enumerate(cards):
    x = 0.7 + index * 4.12
    rect(s, x, 1.95, 3.86, 3.85, WHITE, LINE, True)
    rect(s, x, 1.95, 3.86, 0.08, RED if index < 2 else AMBER)
    text(s, title_, x + 0.28, 2.2, 3.2, 0.32, 15, True, INK)
    text(s, quote, x + 0.28, 2.62, 3.3, 0.85, 11.5, False, INK_2, line=1.5)
    for bi, bullet in enumerate(bullets):
        rect(s, x + 0.30, 3.72 + bi * 0.52, 0.08, 0.08, RED if index < 2 else AMBER, None, True)
        text(s, bullet, x + 0.48, 3.64 + bi * 0.52, 3.0, 0.3, 10.5, False, MUTED)
rect(s, 0.7, 6.1, 11.95, 0.62, RED_SOFT, None, True)
text(s, "结果：统一进度  ×  个体差异  ×  内容不可核验  →  需要一条“有证据、会调整”的学习闭环", 0.95, 6.24, 11.5, 0.32, 12.5, True, RED_DARK)
footer(s, 2)

# ------------------------------------------------------------------ slide 3
s = slide()
header(s, "02", "方案", "一条可执行的学习闭环，而不是一次性推荐")
steps = [
    ("画像", "目标 / 薄弱点\n偏好 / 时间", TEAL_SOFT, TEAL),
    ("证据", "课程 chunk\n来源 / 许可证", BLUE_SOFT, BLUE),
    ("规划", "资源类型\n难度 / 顺序", AMBER_SOFT, AMBER),
    ("生成", "讲义 / 导图\n题库 / 代码", RED_SOFT, RED),
    ("门禁", "质量五维\n安全四类", TEAL_SOFT, TEAL),
    ("路径", "3 天任务\n测验后重规划", BLUE_SOFT, BLUE),
]
for index, (title_, detail, fill, color) in enumerate(steps):
    x = 0.7 + index * 2.02
    rect(s, x, 2.0, 1.72, 1.55, fill, None, True)
    text(s, f"0{index + 1}", x + 0.16, 2.12, 0.5, 0.25, 9, True, color)
    text(s, title_, x + 0.16, 2.4, 1.4, 0.38, 17, True, INK)
    text(s, detail, x + 0.16, 2.88, 1.45, 0.55, 9.5, False, INK_2, line=1.35)
    if index < 5:
        text(s, "→", x + 1.74, 2.62, 0.28, 0.3, 18, True, GRAY_LIGHT)
rect(s, 0.7, 3.95, 5.7, 2.05, WHITE, LINE, True)
text(s, "输入", 0.95, 4.12, 1.0, 0.25, 10, True, MUTED)
text(s, "一句话 + 一门课 + 一份画像", 0.95, 4.42, 5.2, 0.4, 16, True, INK)
text(s, "例：我递归总是绕晕，想要二叉树遍历的图解讲义，最好带 3 道题。", 0.95, 4.9, 5.2, 0.75, 10.5, False, INK_2, line=1.5)
rect(s, 6.95, 3.95, 5.7, 2.05, WHITE, LINE, True)
text(s, "输出", 7.2, 4.12, 1.0, 0.25, 10, True, MUTED)
text(s, "4 类资产 + 可追溯证据", 7.2, 4.42, 5.2, 0.4, 16, True, INK)
text(s, "讲义 / 知识图谱 / 查漏练习 / 学习路径；每条结论带 [E#]，可点回来源。", 7.2, 4.9, 5.2, 0.75, 10.5, False, INK_2, line=1.5)
text(s, "闭环关键：测验结果回流画像与路径，而不是“推荐完就结束”。", 0.7, 6.35, 11.9, 0.35, 12.5, True, RED_DARK)
footer(s, 3)

# ------------------------------------------------------------------ slide 4
s = slide()
header(s, "03", "平台", "为什么把编排放到华为云 AgentArts")
caps = [
    ("智能体 + 工作流", "主控 + 7 角色、12 节点；并行、条件、版本。", TEAL),
    ("平台知识库", "pangu_embedding + pangu_rerank；引用与命中测试开箱即用。", BLUE),
    ("MCP 工具", "工具自描述、自动发现；既有能力低侵入复用。", AMBER),
    ("发布 + 评测", "API/Web 发布；评测集、运行记录、调用日志。", RED),
]
for index, (title_, desc, color) in enumerate(caps):
    y = 1.9 + index * 1.05
    rect(s, 0.7, y, 5.85, 0.88, WHITE, LINE, True)
    rect(s, 0.7, y, 0.08, 0.88, color)
    text(s, title_, 0.95, y + 0.12, 2.1, 0.3, 12.5, True, INK)
    text(s, desc, 2.85, y + 0.13, 3.5, 0.55, 9.5, False, MUTED, line=1.35)
rows = [
    ("智能体/工作流", "代码即配置，改一次发版", "画布配置，改完即试运行"),
    ("知识库", "自建 Chroma/分段/引用", "平台知识库 + 命中测试 + 引用"),
    ("工具", "内部 HTTP，平台看不见", "MCP 自动发现，调用可追踪"),
    ("发布/评测", "自己写 API/鉴权/脚本", "平台发布 + 评测集"),
]
rect(s, 6.9, 1.9, 5.75, 0.42, DARK, None, True)
text(s, "维度", 7.05, 1.97, 1.3, 0.25, 9, True, WHITE)
text(s, "自写编排", 8.45, 1.97, 1.9, 0.25, 9, True, WHITE)
text(s, "AgentArts", 10.45, 1.97, 2.0, 0.25, 9, True, WHITE)
for index, (name, old, new) in enumerate(rows):
    y = 2.36 + index * 0.76
    rect(s, 6.9, y, 5.75, 0.68, WHITE, LINE)
    text(s, name, 7.05, y + 0.19, 1.3, 0.3, 9.5, True, INK)
    text(s, old, 8.45, y + 0.16, 1.95, 0.4, 8.5, False, MUTED, line=1.2)
    text(s, new, 10.45, y + 0.12, 2.05, 0.45, 8.5, True, TEAL, line=1.2)
text(s, "旁路适配：原项目 0 行修改，既有 RAG/生成/评估通过 MCP 复用。", 6.9, 5.55, 5.75, 0.55, 11.5, True, RED_DARK)
rect(s, 6.9, 6.1, 5.75, 0.5, RED_SOFT, None, True)
text(s, "平台负责编排/知识/发布/评测；领域能力仍归 EduPath。", 7.05, 6.22, 5.45, 0.3, 9.5, True, RED_DARK)
footer(s, 4)

# ------------------------------------------------------------------ slide 5
s = slide()
header(s, "04", "编制", "8 个角色，各自对一件事负责")
agents = [
    ("知径学伴", "主控 / Router", "意图识别、路由、汇总", "DeepSeek-V4-Flash"),
    ("画像师", "Profile", "目标、薄弱点、偏好、节奏", "temp 0.1"),
    ("检索员", "Retriever", "课程证据、阈值、拒答", "pangu_embedding"),
    ("教学设计", "Planner", "资源组合、难度、时长", "temp 0.3"),
    ("内容生成", "Generator", "讲义 / 导图 / 题库 / 代码", "temp 0.4"),
    ("辅导教练", "Tutor", "苏格拉底 / 分步 / 提示", "temp 0.3"),
    ("评估师", "Evaluator", "五维质量、测验分析", "temp 0.0"),
    ("安全员", "Safety", "诚信 / 隐私 / 事实 / 价值观", "temp 0.0"),
]
for index, (name, role, duty, model) in enumerate(agents):
    col, row = index % 4, index // 4
    x = 0.7 + col * 3.12
    y = 1.9 + row * 2.15
    rect(s, x, y, 2.9, 1.9, WHITE, LINE, True)
    rect(s, x, y, 0.07, 1.9, RED if index == 0 else TEAL if index < 4 else BLUE)
    text(s, name, x + 0.22, y + 0.16, 2.1, 0.32, 13.5, True, INK)
    text(s, role, x + 0.22, y + 0.5, 2.2, 0.24, 8.5, True, MUTED)
    text(s, duty, x + 0.22, y + 0.82, 2.4, 0.55, 9.5, False, INK_2, line=1.4)
    text(s, model, x + 0.22, y + 1.5, 2.4, 0.24, 8, False, FAINT)
text(s, "不是为 multi-agent 而 multi-agent：职责与失败边界不同，才能独立重试、独立评测、独立换模型。", 0.7, 6.28, 11.9, 0.4, 11.5, True, RED_DARK)
footer(s, 5)

# ------------------------------------------------------------------ slide 6
s = slide()
header(s, "05", "编排", "12 节点工作流：并行、门禁、可回退")
top = [
    ("开始", "消息 / 课程"),
    ("意图识别", "主控 Agent"),
    ("课程确认", "校验 / 回退"),
    ("画像 + 检索", "并行"),
    ("资源规划", "类型 / 难度"),
]
for index, (title_, detail) in enumerate(top):
    x = 0.7 + index * 2.42
    rect(s, x, 1.95, 2.06, 0.95, WHITE, LINE, True)
    text(s, title_, x + 0.14, 2.05, 1.8, 0.3, 11, True, INK)
    text(s, detail, x + 0.14, 2.38, 1.8, 0.28, 8.5, False, MUTED)
    if index < 4:
        text(s, "→", x + 2.07, 2.25, 0.3, 0.3, 16, True, GRAY_LIGHT)
gen = ["生成讲义", "生成导图", "生成题库"]
for index, label in enumerate(gen):
    x = 0.7 + index * 1.55
    rect(s, x, 3.35, 1.35, 0.6, RED_SOFT, RGBColor(0xF0, 0xD4, 0xD9), True)
    text(s, label, x, 3.48, 1.35, 0.25, 9.5, True, RED_DARK, PP_ALIGN.CENTER)
text(s, "三路并行生成", 5.35, 3.42, 1.8, 0.3, 9.5, True, MUTED)
rect(s, 7.2, 3.3, 1.55, 0.7, AMBER_SOFT, None, True)
text(s, "质量门禁", 7.2, 3.43, 1.55, 0.25, 10, True, AMBER, PP_ALIGN.CENTER)
text(s, "→", 8.8, 3.45, 0.3, 0.3, 16, True, GRAY_LIGHT)
rect(s, 9.15, 3.3, 1.55, 0.7, TEAL_SOFT, None, True)
text(s, "安全审核", 9.15, 3.43, 1.55, 0.25, 10, True, TEAL, PP_ALIGN.CENTER)
text(s, "→", 10.75, 3.45, 0.3, 0.3, 16, True, GRAY_LIGHT)
rect(s, 11.1, 3.3, 1.55, 0.7, WHITE, LINE, True)
text(s, "发布 / 落库", 11.1, 3.43, 1.55, 0.25, 10, True, INK, PP_ALIGN.CENTER)
rect(s, 7.2, 4.35, 5.45, 0.62, BLUE_SOFT, None, True)
text(s, "路径规划：3 天 × 30 分钟，按测验结果重规划", 7.35, 4.48, 5.2, 0.3, 10, True, BLUE)
text(s, "质量 / 安全通过后才进入发布", 8.35, 4.98, 3.6, 0.3, 9, False, MUTED)
rect(s, 0.7, 5.55, 11.95, 0.95, WHITE, LINE, True)
text(s, "测验分支", 0.95, 5.68, 1.2, 0.3, 10, True, RED)
text(s, "提交作答 → 评估 Agent（掌握度 / 错因）→ 插入补救或调整难度 → 路径重规划 → 写回画像", 2.1, 5.66, 10.3, 0.35, 10.5, False, INK_2)
text(s, "失败可回退：检索 0 命中拒答；生成失败重试 1 次；质量不通过返回草稿；安全高危直接拦截。", 0.7, 6.68, 11.9, 0.32, 9.5, True, MUTED)
footer(s, 6)

# ------------------------------------------------------------------ slide 7
s = slide()
header(s, "06", "知识", "知识接入：每条结论都能回到课程 chunk")
stats = [("2", "门课程"), ("108", "篇公开语料"), ("1,200+", "语义切片")]
for index, (num, label) in enumerate(stats):
    x = 0.7 + index * 1.85
    text(s, num, x, 1.9, 1.55, 0.55, 26, True, INK)
    text(s, label, x, 2.48, 1.55, 0.25, 9.5, False, MUTED)
pipeline = ["授权原文", "清洗 / 分段", "元数据", "平台知识库"]
for index, label in enumerate(pipeline):
    y = 2.95 + index * 0.62
    rect(s, 0.7, y, 2.5, 0.46, WHITE, LINE, True)
    text(s, f"0{index + 1}", 0.84, y + 0.1, 0.35, 0.22, 8, True, RED)
    text(s, label, 1.2, y + 0.08, 1.8, 0.28, 9.5, True, INK)
    if index < 3:
        text(s, "↓", 1.05, y + 0.45, 0.3, 0.2, 10, True, GRAY_LIGHT)
text(s, "分段：标题层级 3 / 约 500 字 / 重叠 50–80 字", 0.7, 5.55, 4.9, 0.3, 9, False, MUTED)
rect(s, 0.7, 5.95, 4.95, 0.72, TEAL_SOFT, None, True)
text(s, "负例：红黑树删除 → 0 命中 → 拒答", 0.88, 6.08, 4.6, 0.28, 10.5, True, TEAL)
text(s, "该拒答时拒答，比多答一题更能证明“证据约束”。", 0.88, 6.36, 4.6, 0.24, 8.5, False, RGBColor(0x4C, 0x80, 0x7A))
rect(s, 5.95, 1.85, 6.7, 4.95, WHITE, LINE, True)
text(s, "讲义片段 · 带引用", 6.15, 2.0, 3.2, 0.25, 8.5, True, MUTED)
text(s, "递归三要素（节选）", 6.15, 2.28, 6.2, 0.4, 17, True, INK)
for index, line in enumerate([
    "1. 终止条件：节点为空时返回，空节点是递归出口，不是异常 [E2]。",
    "2. 本层动作：把当前节点的值加入结果列表。",
    "3. 下一层：分别递归左、右子树；递归深度等于树高 [E3]。",
]):
    text(s, line, 6.15, 2.8 + index * 0.34, 6.2, 0.28, 9.5, False, INK_2, line=1.35)
rect(s, 6.15, 3.9, 6.25, 1.2, DARK_2, None, True)
code = ("def inorder(root):\n"
        "    if root is None:          # 1. 出口\n"
        "        return []\n"
        "    return inorder(root.left) + [root.val] + inorder(root.right)")
text(s, code, 6.32, 4.03, 5.95, 0.98, 8.2, False, RGBColor(0xE6, 0xEE, 0xF2), line=1.35, font="Consolas")
text(s, "证据侧栏（可点回 chunk）", 6.15, 5.24, 4.5, 0.24, 8.5, True, MUTED)
evidence_rows = [
    ("E1", "前序、中序、后序与层序遍历", "0.94", "OpenDSA / MIT"),
    ("E2", "二叉树术语、性质与结构分类", "0.94", "OpenDSA / MIT"),
    ("E3", "DFS：时间戳、边分类与递归栈", "0.90", "cp-algorithms / CC BY-SA"),
]
for index, (eid, title_, score, source) in enumerate(evidence_rows):
    y = 5.52 + index * 0.36
    rect(s, 6.15, y, 0.34, 0.22, DARK, None, True)
    text(s, eid, 6.15, y + 0.03, 0.34, 0.16, 7, True, WHITE, PP_ALIGN.CENTER)
    text(s, title_, 6.58, y - 0.01, 3.65, 0.22, 8.3, True, INK)
    text(s, score, 10.3, y - 0.01, 0.45, 0.22, 7.5, True, TEAL)
    text(s, source, 10.8, y - 0.01, 1.7, 0.22, 6.8, False, FAINT)
text(s, "每条结论都能点回 chunk；弱证据拒答，而不是让模型补常识。", 6.15, 6.5, 6.3, 0.24, 7.8, False, MUTED)
footer(s, 7)

# ------------------------------------------------------------------ slide 8
s = slide()
header(s, "07", "工具", "10 个 MCP 工具，把既有能力变成可编排动作")
tool_names = [
    "画像读取", "课程检索", "资源规划", "资源生成", "路径规划",
    "证据辅导", "出题", "测验分析", "安全审核", "课程大纲",
]
for index, name in enumerate(tool_names):
    col, row = index % 2, index // 2
    x = 0.7 + col * 2.55
    y = 1.92 + row * 0.82
    rect(s, x, y, 2.35, 0.66, WHITE, LINE, True)
    text(s, f"{index + 1:02d}", x + 0.12, y + 0.18, 0.35, 0.22, 8, True, RED)
    text(s, name, x + 0.52, y + 0.15, 1.7, 0.3, 10.5, True, INK)
    text(s, "edupath_*", x + 0.52, y + 0.4, 1.6, 0.2, 7.5, False, FAINT)
rect(s, 6.25, 1.9, 6.4, 2.65, DARK_2, None, True)
text(s, "tools/call  ·  JSON-RPC 2.0", 6.45, 2.05, 5.8, 0.25, 8.5, True, RGBColor(0x9F, 0xC7, 0xC0))
code = ('{\n'
        '  "jsonrpc": "2.0", "id": 3,\n'
        '  "method": "tools/call",\n'
        '  "params": {\n'
        '    "name": "edupath_search_knowledge",\n'
        '    "arguments": {"query": "二叉树 中序",\n'
        '                  "course_id": 1, "top_k": 4}\n'
        '  }\n'
        '}')
text(s, code, 6.45, 2.38, 5.9, 1.95, 8.5, False, RGBColor(0xE6, 0xEE, 0xF2), line=1.35, font="Consolas")
text(s, "stdio  ·  HTTP JSON-RPC  ·  HTTP + SSE", 6.25, 4.65, 6.3, 0.3, 9, True, MUTED)
rect(s, 6.25, 5.1, 6.4, 1.55, RED_SOFT, None, True)
text(s, "工具发现 → 参数绑定 → 调用日志", 6.45, 5.25, 6.0, 0.3, 11.5, True, RED_DARK)
text(s, "评委能看到 Agent 真的调了工具，而不是“写了一段像工具调用的话”。", 6.45, 5.62, 5.9, 0.55, 9.5, False, RGBColor(0x8F, 0x3A, 0x4A), line=1.4)
text(s, "原项目 0 行修改：MCP 把既有 HTTP 能力注册为工具，平台负责编排。", 6.25, 6.0, 6.4, 0.35, 8.5, True, RED_DARK)
footer(s, 8)

# ------------------------------------------------------------------ slide 9
s = slide()
header(s, "08", "Demo", "一个学生、一门课、一次完整学习任务")
shots = [
    ("demo-home.png", "01  画像与任务", "一句话输入，进入工作流"),
    ("demo-quiz.png", "02  练习与评估", "3 题查漏，错因回流路径"),
    ("demo-path.png", "03  三日路径", "按薄弱点排任务与时长"),
]
for index, (name, title_, caption) in enumerate(shots):
    x = 0.7 + index * 4.12
    picture(s, EVIDENCE / name, x, 1.85, 3.86, 2.55, border=True)
    text(s, title_, x, 4.52, 3.86, 0.3, 11.5, True, INK)
    text(s, caption, x, 4.85, 3.86, 0.3, 9, False, MUTED)
rect(s, 0.7, 5.45, 11.95, 0.8, DARK, None, True)
text(s, "9 步轨迹   ·   5 条课程证据   ·   3 类资源   ·   88 分质量门禁   ·   4 项安全通过", 0.95, 5.62, 11.5, 0.35, 12.5, True, WHITE, PP_ALIGN.CENTER)
text(s, "截图来自 offline fixture 彩排，页面明确标注“本地演示 · 未上云”；正式提交前替换为 AgentArts 端点截图。", 0.7, 6.5, 11.9, 0.3, 8.5, False, FAINT)
footer(s, 9)

# ------------------------------------------------------------------ slide 10
s = slide()
header(s, "09", "质量", "质量与安全：生成不等于发布")
rect(s, 0.7, 1.9, 3.1, 3.65, DARK, None, True)
text(s, "88", 0.9, 2.15, 2.7, 1.0, 52, True, WHITE)
text(s, "/ 100", 2.55, 2.65, 0.9, 0.3, 12, True, RGBColor(0x9F, 0xB0, 0xBC))
text(s, "A 级 · 门禁通过", 0.9, 3.25, 2.7, 0.3, 12, True, RGBColor(0x8F, 0xC9, 0xC1))
text(s, "五维加权：课程一致性 30%\n学情匹配 25% / 结构 20%\n可练习 15% / 表达 10%", 0.9, 3.75, 2.7, 1.2, 9.5, False, RGBColor(0xCF, 0xD8, 0xDE), line=1.5)
dims = [("课程一致性", 92), ("学情匹配", 90), ("结构完整性", 86), ("可练习性", 84), ("表达清晰度", 86)]
for index, (name, score) in enumerate(dims):
    y = 1.98 + index * 0.62
    text(s, name, 4.15, y, 1.15, 0.25, 9, False, INK_2)
    rect(s, 5.35, y + 0.08, 2.2, 0.12, RGBColor(0xEC, 0xE6, 0xDB), None, True)
    rect(s, 5.35, y + 0.08, 2.2 * score / 100, 0.12, TEAL, None, True)
    text(s, str(score), 7.65, y, 0.45, 0.25, 8.5, True, TEAL)
rect(s, 8.35, 1.9, 4.3, 3.65, WHITE, LINE, True)
text(s, "安全审核", 8.55, 2.05, 3.8, 0.3, 11.5, True, INK)
checks = ["学术诚信：不代写/代考", "隐私：不重复敏感信息", "事实性：引用必须存在", "价值观：拦截高危内容"]
for index, item in enumerate(checks):
    y = 2.48 + index * 0.58
    rect(s, 8.55, y + 0.02, 0.16, 0.16, TEAL, None, True)
    text(s, item, 8.82, y - 0.03, 3.6, 0.3, 9.5, False, INK_2)
rect(s, 8.55, 4.85, 3.85, 0.5, TEAL_SOFT, None, True)
text(s, "高危直接拦截，不进入发布", 8.7, 4.96, 3.6, 0.26, 9.5, True, TEAL)
text(s, "工程验证 82 / 56 / 29 / 2（AI / Java / 前端 / E2E）——证明工程流程，不宣称提分。", 0.7, 5.85, 11.9, 0.3, 9, False, MUTED)
rect(s, 0.7, 6.3, 11.95, 0.55, AMBER_SOFT, None, True)
text(s, "发布前校验引用：模型写了 [E1]，E1 必须真的在检索结果里；否则降级草稿。", 0.95, 6.42, 11.5, 0.3, 10, True, AMBER)
footer(s, 10)

# ------------------------------------------------------------------ slide 11
s = slide()
header(s, "10", "价值", "应用价值：按课程复制，而不是按学校重做")
models = [
    ("课程包", "给院系/课程团队", "语料 + 知识点 + 工作流 + 评测集", TEAL_SOFT, TEAL),
    ("教师工作台", "给任课教师", "分层讲义、练习、班级洞察", BLUE_SOFT, BLUE),
    ("SaaS 订阅", "给高校/职教", "按课程/班级计费，持续更新", RED_SOFT, RED),
]
for index, (title_, target, desc, fill, color) in enumerate(models):
    x = 0.7 + index * 4.12
    rect(s, x, 1.9, 3.86, 1.75, fill, None, True)
    text(s, title_, x + 0.25, 2.05, 3.2, 0.35, 15, True, INK)
    text(s, target, x + 0.25, 2.42, 3.2, 0.25, 9, True, color)
    text(s, desc, x + 0.25, 2.75, 3.35, 0.65, 9.5, False, INK_2, line=1.4)
text(s, "复制路径", 0.7, 4.0, 1.3, 0.3, 11, True, MUTED)
copy_steps = ["换语料", "标知识点", "配评测集", "跑工作流"]
for index, label in enumerate(copy_steps):
    x = 1.95 + index * 2.55
    rect(s, x, 3.95, 1.9, 0.55, WHITE, LINE, True)
    text(s, label, x, 4.08, 1.9, 0.26, 10, True, INK, PP_ALIGN.CENTER)
    if index < 3:
        text(s, "→", x + 1.92, 4.06, 0.3, 0.3, 14, True, GRAY_LIGHT)
metrics = [
    ("45 → 15 分钟", "教师备课（目标，待试点）"),
    ("≥ 95%", "任务成功率（目标）"),
    ("100%", "证据覆盖 / 引用可追溯"),
    ("100%", "高危样本拦截（评测集）"),
]
for index, (value, label) in enumerate(metrics):
    x = 0.7 + index * 3.0
    text(s, value, x, 5.0, 2.9, 0.45, 20, True, RED if index == 0 else INK)
    text(s, label, x, 5.5, 2.9, 0.3, 8.5, False, MUTED)
text(s, "不宣称提升学习成绩：没有对照实验；下一步做班级试点，记录完成率与教师备课时间。", 0.7, 6.35, 11.9, 0.35, 9.5, True, RED_DARK)
footer(s, 11)

# ------------------------------------------------------------------ slide 12
s = slide(DARK)
rect(s, 0, 0, 0.16, 7.5, RED)
text(s, "知径 EduPath", 0.8, 0.75, 6.0, 0.6, 30, True, WHITE)
text(s, "让学习有据可循。", 0.8, 1.45, 8.0, 0.55, 22, False, RGBColor(0x8F, 0xC9, 0xC1))
text(s, "基于华为云 AgentArts 的课程学习多智能体", 0.8, 2.08, 8.0, 0.35, 11, False, RGBColor(0xB9, 0xC4, 0xCC))
cols = [
    ("已完成", ["本地 Demo / MCP / 9 步流水线", "信息表 / 12 页 PPT / 5 张截图", "16 项自检 + 14 项单测"], RGBColor(0x8F, 0xC9, 0xC1)),
    ("待完成", ["AgentArts 真实端点联调", "3.5 分钟实机视频", "24 条平台评测"], RGBColor(0xE8, 0xB4, 0x8A)),
    ("边界", ["2 门课，不宣称全学科", "不宣称提分，无对照实验", "fixture 截图已标注未上云"], RGBColor(0xC9, 0xA3, 0xAC)),
]
for index, (title_, items, color) in enumerate(cols):
    x = 0.8 + index * 4.1
    text(s, title_, x, 3.0, 3.6, 0.3, 10, True, color)
    for item_index, item in enumerate(items):
        text(s, f"· {item}", x, 3.42 + item_index * 0.42, 3.6, 0.3, 9.5, False, RGBColor(0xD5, 0xDD, 0xE3))
text(s, "github.com/Richard-Lee0419/-EduPath  ·  branch: huawei-ict-2026", 0.8, 5.95, 11.5, 0.3, 10, True, WHITE)
text(s, "谢谢，欢迎提问。", 0.8, 6.35, 5.0, 0.35, 13, False, RGBColor(0x8F, 0xC9, 0xC1))
text(s, "注：本页“已完成/待完成”为依据当前仓库状态的如实标注，不构成上线或效果承诺。", 0.8, 6.85, 11.8, 0.25, 7.5, False, RGBColor(0x6F, 0x7F, 0x8B))

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"written: {OUT}")
print(f"slides: {len(prs.slides)}")
