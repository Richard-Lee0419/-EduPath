from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


SERVICE_ROOT = Path(__file__).resolve().parents[1]
if str(SERVICE_ROOT) not in sys.path:
    sys.path.insert(0, str(SERVICE_ROOT))


def main() -> int:
    parser = argparse.ArgumentParser(description="Inventory EduPath courseware before corpus cataloging.")
    parser.add_argument("--source-dir", default=str(SERVICE_ROOT / "data/course_corpus/source_documents"))
    parser.add_argument("--report", default=str(SERVICE_ROOT / "data/course_corpus/material_inventory.json"))
    args = parser.parse_args()

    source_root = Path(args.source_dir).resolve()
    files = [path for path in source_root.rglob("*") if path.is_file() and path.name != ".gitkeep"]
    inventory: list[dict[str, Any]] = []
    for path in sorted(files, key=lambda item: item.as_posix().lower()):
        relative_path = path.relative_to(source_root).as_posix()
        try:
            item = _inventory_file(path)
            item.update({"path": relative_path, "bytes": path.stat().st_size, "status": "parsed"})
        except Exception as exception:
            item = {
                "path": relative_path,
                "bytes": path.stat().st_size,
                "status": "failed",
                "error": f"{type(exception).__name__}: {exception}",
                "pages_or_slides": 0,
                "text_characters": 0,
                "empty_units": 0,
                "sample": "",
                "headings": [],
            }
        inventory.append(item)

    by_folder = []
    for folder, rows in _group_by_folder(inventory).items():
        by_folder.append(
            {
                "folder": folder,
                "files": len(rows),
                "parsed": sum(row["status"] == "parsed" for row in rows),
                "failed": sum(row["status"] == "failed" for row in rows),
                "pages_or_slides": sum(int(row["pages_or_slides"]) for row in rows),
                "text_characters": sum(int(row["text_characters"]) for row in rows),
            }
        )
    report = {
        "version": 1,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_root": "data/course_corpus/source_documents",
        "summary": {
            "files": len(inventory),
            "parsed": sum(row["status"] == "parsed" for row in inventory),
            "failed": sum(row["status"] == "failed" for row in inventory),
            "pages_or_slides": sum(int(row["pages_or_slides"]) for row in inventory),
            "text_characters": sum(int(row["text_characters"]) for row in inventory),
            "extensions": dict(Counter(Path(row["path"]).suffix.lower() for row in inventory)),
        },
        "folders": by_folder,
        "files": inventory,
    }
    report_path = Path(args.report)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"ok": report["summary"]["failed"] == 0, "report": args.report, **report["summary"]}, ensure_ascii=False))
    return 0 if report["summary"]["failed"] == 0 else 1


def _inventory_file(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _inventory_pdf(path)
    if suffix == ".pptx":
        return _inventory_pptx(path)
    if suffix == ".docx":
        return _inventory_docx(path)
    raise ValueError(f"unsupported extension: {suffix}")


def _inventory_pdf(path: Path) -> dict[str, Any]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    units = [page.extract_text() or "" for page in reader.pages]
    return _text_inventory("pdf", units)


def _inventory_pptx(path: Path) -> dict[str, Any]:
    from pptx import Presentation

    presentation = Presentation(path)
    units = []
    for slide in presentation.slides:
        units.append("\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()))
    return _text_inventory("pptx", units)


def _inventory_docx(path: Path) -> dict[str, Any]:
    from docx import Document

    document = Document(path)
    units = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    return _text_inventory("docx", units)


def _text_inventory(kind: str, units: list[str]) -> dict[str, Any]:
    normalized = [_normalize_text(text) for text in units]
    combined = "\n\n".join(text for text in normalized if text)
    headings: list[str] = []
    for text in normalized:
        for line in text.splitlines()[:4]:
            candidate = " ".join(line.split())
            if 4 <= len(candidate) <= 140 and candidate not in headings:
                headings.append(candidate)
                break
        if len(headings) >= 12:
            break
    return {
        "kind": kind,
        "pages_or_slides": len(units),
        "text_characters": len(combined),
        "empty_units": sum(not text for text in normalized),
        "average_characters_per_unit": round(len(combined) / len(units), 1) if units else 0.0,
        "sample": combined[:1200],
        "headings": headings,
    }


def _normalize_text(text: str) -> str:
    return "\n".join(line.strip() for line in text.replace("\r", "\n").split("\n") if line.strip())


def _group_by_folder(rows: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        folder = str(row["path"]).split("/", 1)[0]
        grouped.setdefault(folder, []).append(row)
    return grouped


if __name__ == "__main__":
    raise SystemExit(main())
