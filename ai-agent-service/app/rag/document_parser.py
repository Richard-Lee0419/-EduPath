from __future__ import annotations

from io import BytesIO


class DocumentParseError(ValueError):
    """Raised when an uploaded course document cannot yield readable text."""


class DocumentTooLargeError(DocumentParseError):
    def __init__(self, actual_bytes: int, max_bytes: int) -> None:
        self.actual_bytes = actual_bytes
        self.max_bytes = max_bytes
        super().__init__(f"文档大小超过限制: {max_bytes} bytes")


def parse_document_bytes(
    raw: bytes,
    *,
    filename: str,
    content_type: str | None = None,
    max_bytes: int | None = None,
) -> str:
    """Parse a supported course document into normalized plain text."""

    if max_bytes is not None and len(raw) > max_bytes:
        raise DocumentTooLargeError(len(raw), max_bytes)

    lowered = filename.lower()
    media_type = (content_type or "").lower()
    if lowered.endswith(".pdf") or "pdf" in media_type:
        text = _parse_pdf(raw)
    elif lowered.endswith(".docx") or "word" in media_type:
        text = _parse_docx(raw)
    elif lowered.endswith(".pptx") or "presentation" in media_type:
        text = _parse_pptx(raw)
    else:
        text = _decode_text(raw)

    normalized = _normalize_text(text)
    if not normalized:
        raise DocumentParseError("文档未解析到有效文本，无法写入知识库索引")
    return normalized


def _parse_pdf(raw: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(raw))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exception:
        raise DocumentParseError("PDF 文档解析失败或未包含可提取文本") from exception


def _parse_docx(raw: bytes) -> str:
    try:
        from docx import Document

        document = Document(BytesIO(raw))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        return "\n\n".join(paragraphs)
    except Exception as exception:
        raise DocumentParseError("DOCX 文档解析失败或未包含可提取文本") from exception


def _parse_pptx(raw: bytes) -> str:
    try:
        from pptx import Presentation

        presentation = Presentation(BytesIO(raw))
        slides: list[str] = []
        for index, slide in enumerate(presentation.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            if texts:
                slides.append(f"Slide {index}\n" + "\n".join(texts))
        return "\n\n".join(slides)
    except Exception as exception:
        raise DocumentParseError("PPTX 文档解析失败或未包含可提取文本") from exception


def _decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    try:
        return raw.decode("latin-1")
    except UnicodeDecodeError as exception:
        raise DocumentParseError("文本编码无法识别") from exception


def _normalize_text(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    compact: list[str] = []
    previous_blank = False
    for line in lines:
        blank = not line.strip()
        if blank and previous_blank:
            continue
        compact.append(line)
        previous_blank = blank
    return "\n".join(compact).strip()
